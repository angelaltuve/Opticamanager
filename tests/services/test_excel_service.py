"""Tests for the Excel export service (excel_service)."""

from __future__ import annotations

from pathlib import Path

import openpyxl

from opticamanager.services import excel_service


class TestExportReportToExcel:
    def test_exporta_workbook_con_titulo_y_subtitulo(self, tmp_path: Path) -> None:
        """Title/subtitle are merged on top; sheet inherits truncated title."""
        out = tmp_path / "reporte.xlsx"

        excel_service.export_report_to_excel(
            out,
            headers=["Producto", "Total"],
            data=[["Lentes", "Bs. 1.234,56"]],
            title="Reporte de Ventas",
            subtitle="Enero 2026",
        )

        assert out.is_file()
        wb = openpyxl.load_workbook(out)
        ws = wb.active
        assert ws.title == "Reporte de Ventas"
        # (1,1) = title, (2,1) = subtitle, (3,1) = blank, row 4 = headers
        assert ws.cell(row=1, column=1).value == "Reporte de Ventas"
        assert ws.cell(row=2, column=1).value == "Enero 2026"
        assert ws.cell(row=4, column=1).value == "Producto"
        assert ws.cell(row=4, column=2).value == "Total"
        assert ws.cell(row=5, column=1).value == "Lentes"
        # Header formatting: dark fill + white bold font
        assert ws.cell(row=4, column=1).font.bold is True
        assert ws.cell(row=4, column=1).font.color.rgb == "00FFFFFF"
        # Freeze panes below the header row
        assert ws.freeze_panes == "A5"

    def test_exporta_sin_titulo_usa_nombre_generico(self, tmp_path: Path) -> None:
        out = tmp_path / "simple.xlsx"

        excel_service.export_report_to_excel(
            out,
            headers=["A", "B"],
            data=[["x", "y"]],
        )

        wb = openpyxl.load_workbook(out)
        assert wb.active.title == "Reporte"
        assert wb.active.cell(row=1, column=1).value == "A"  # headers start row 1
        assert wb.active.cell(row=2, column=1).value == "x"
        assert wb.active.freeze_panes == "A2"

    def test_titulo_largo_truncado_a_31_caracteres(self, tmp_path: Path) -> None:
        out = tmp_path / "truncado.xlsx"
        long_title = "R" * 60

        excel_service.export_report_to_excel(out, ["A"], [["v"]], title=long_title)

        wb = openpyxl.load_workbook(out)
        assert wb.active.title == "R" * 31

    def test_valores_monetarios_se_convierten_a_numero(self, tmp_path: Path) -> None:
        """Headers matching CURRENCY_HINTS get float values + #,##0.00 format."""
        out = tmp_path / "moneda.xlsx"

        excel_service.export_report_to_excel(
            out,
            headers=["Producto", "Precio", "Total"],
            data=[["Lentes", "Bs. 36,50", "$ 25,00"]],
        )

        wb = openpyxl.load_workbook(out)
        ws = wb.active
        # Venezuelan decimal policy: comma = decimal separator, dot = thousands
        assert ws.cell(row=2, column=2).value == 36.5  # "Bs. 36,50" → 36.5
        assert ws.cell(row=2, column=2).number_format == "#,##0.00"
        assert ws.cell(row=2, column=3).value == 25.0  # "$ 25,00" → 25.0
        assert ws.cell(row=2, column=3).number_format == "#,##0.00"

    def test_parsing_numerico_venezolano(self, tmp_path: Path) -> None:
        """Comma = decimal separator, dot = thousands separator."""
        out = tmp_path / "parse.xlsx"

        excel_service.export_report_to_excel(
            out,
            headers=["Total"],
            data=[
                ["1.234,56"],  # thousands + decimals
                ["36,50"],  # decimals only
                ["0,50"],
                ["1.000"],  # thousands only
                ["-1.234,56"],  # negative
                ["12.345,67"],
                ["123"],  # plain integer
                ["1.5"],  # plain dot float
                ["N/D"],  # text passes through
                ["—"],  # placeholder passes through
            ],
        )

        wb = openpyxl.load_workbook(out)
        ws = wb.active
        assert ws.cell(row=2, column=1).value == 1234.56
        assert ws.cell(row=3, column=1).value == 36.5
        assert ws.cell(row=4, column=1).value == 0.5
        assert ws.cell(row=5, column=1).value == 1000.0
        assert ws.cell(row=6, column=1).value == -1234.56
        assert ws.cell(row=7, column=1).value == 12345.67
        assert ws.cell(row=8, column=1).value == 123
        assert ws.cell(row=9, column=1).value == 1.5
        assert ws.cell(row=10, column=1).value == "N/D"
        assert ws.cell(row=11, column=1).value == "—"
        # Currency in the CURRENCY_HINTS columns
        assert ws.cell(row=2, column=1).number_format == "#,##0.00"

    def test_valores_enteros_se_convierten_a_numero_sin_decimales(
        self, tmp_path: Path
    ) -> None:
        """Headers matching INTEGER_HINTS get #,##0 (no decimals)."""
        out = tmp_path / "enteros.xlsx"

        excel_service.export_report_to_excel(
            out,
            headers=["Producto", "Cant.", "Stock"],
            data=[["Lentes", "12", "45"]],
        )

        wb = openpyxl.load_workbook(out)
        ws = wb.active
        assert ws.cell(row=2, column=2).value == 12
        assert ws.cell(row=2, column=2).number_format == "#,##0"
        assert ws.cell(row=2, column=3).value == 45
        assert ws.cell(row=2, column=3).number_format == "#,##0"

    def test_valores_no_numericos_permanecen_como_texto(self, tmp_path: Path) -> None:
        out = tmp_path / "texto.xlsx"

        excel_service.export_report_to_excel(
            out,
            headers=["Producto", "Nota"],
            data=[["Lentes", "N/D"], ["Monturas", "—"], ["Aros", ""]],
        )

        wb = openpyxl.load_workbook(out)
        ws = wb.active
        assert ws.cell(row=2, column=2).value == "N/D"
        assert ws.cell(row=3, column=2).value == "—"
        # openpyxl normalizes "" cells to None when it re-reads the file
        assert ws.cell(row=4, column=2).value in ("", None)

    def test_data_vacia_escribe_solo_encabezados(self, tmp_path: Path) -> None:
        out = tmp_path / "vacio.xlsx"

        excel_service.export_report_to_excel(out, ["A", "B"], data=[])

        wb = openpyxl.load_workbook(out)
        ws = wb.active
        assert ws.cell(row=1, column=1).value == "A"
        assert ws.cell(row=1, column=2).value == "B"
        assert ws.max_row == 1

    def test_ancho_de_columnas_calculado_y_limitado(self, tmp_path: Path) -> None:
        out = tmp_path / "anchos.xlsx"

        excel_service.export_report_to_excel(
            out,
            headers=["Corta", "MuyLargaHeaderQueExcedeElLimiteDeCuarentaYCinco"],
            data=[["a", "b"]],
        )

        wb = openpyxl.load_workbook(out)
        ws = wb.active
        col_a = ws.column_dimensions["A"].width
        col_b = ws.column_dimensions["B"].width
        assert col_a == 9  # len("Corta") + 4
        assert col_b == 45  # capped at 45

    def test_exporta_y_registra_log(self, tmp_path: Path, caplog) -> None:
        out = tmp_path / "log.xlsx"

        with caplog.at_level("INFO", logger="opticamanager.services.excel_service"):
            excel_service.export_report_to_excel(out, ["A"], [["1"], ["2"]], title="T")

        assert any("Excel exportado" in r.message for r in caplog.records)
